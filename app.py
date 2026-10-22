"""
AI Invoice Generator  -  Streamlit + Gemini Flash
-------------------------------------------------
* Live HTML preview (printable)
* PDF export (ReportLab) and Excel export (openpyxl, with live formulas)
* Optional AI assistant powered by Google Gemini Flash (free tier key)
"""

import base64
import io
import json
import os
import re
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from html import escape as html_escape
from xml.sax.saxutils import escape as xml_escape

import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    Image as RLImage,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------
CURRENCIES = {
    "USD": "$",
    "EUR": "€",
    "GBP": "£",
    "PKR": "Rs ",
    "INR": "Rs ",
    "AED": "AED ",
    "SAR": "SAR ",
    "CAD": "CA$",
    "AUD": "A$",
    "JPY": "¥",
}
DISCOUNT_TYPES = ["None", "Percentage (%)", "Fixed amount"]
DEFAULT_MODEL = "gemini-flash-latest"  # Google-maintained alias -> always the current Flash model
FALLBACK_MODELS = ["gemini-flash-lite-latest", "gemini-2.5-flash"]  # tried if the main model 404s / has no quota
ITEM_COLS = ["Description", "Quantity", "Unit Price"]

FORM_KEYS = [
    "biz_name", "biz_email", "biz_phone", "biz_address", "biz_tax_id",
    "cust_name", "cust_email", "cust_phone", "cust_address",
    "inv_no", "inv_date", "due_date", "currency", "accent",
    "notes", "terms",
    "tax_pct", "discount_type", "discount_val", "shipping",
]


# --------------------------------------------------------------------------
# Money / calculation helpers (Decimal => exact, no floating point drift)
# --------------------------------------------------------------------------
def to_dec(value) -> Decimal:
    """Safely convert anything (None, NaN, str, float) to a non-negative Decimal."""
    try:
        if value is None:
            return Decimal("0")
        num = pd.to_numeric(value, errors="coerce")
        if pd.isna(num):
            return Decimal("0")
        d = Decimal(str(num))
        if not d.is_finite() or d < 0:
            return Decimal("0")
        return d
    except Exception:
        return Decimal("0")


def q2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def fmt_qty(value: Decimal) -> str:
    if value == value.to_integral_value():
        return str(int(value))
    return f"{value:f}".rstrip("0").rstrip(".")


def money(sym: str, value: Decimal) -> str:
    return f"{sym}{value:,.2f}"


def clean_items(df: pd.DataFrame) -> list:
    """Turn the editor dataframe into a list of validated line items."""
    items = []
    if df is None or len(df) == 0:
        return items
    for _, row in df.iterrows():
        desc = row.get("Description", "")
        desc = "" if desc is None or (isinstance(desc, float) and pd.isna(desc)) else str(desc).strip()
        qty = to_dec(row.get("Quantity"))
        price = to_dec(row.get("Unit Price"))
        amount = q2(qty * price)
        if not desc and amount == 0:
            continue  # ignore completely empty rows
        items.append({"desc": desc or "(no description)", "qty": qty, "price": q2(price), "amount": amount})
    return items


def compute_totals(items, tax_pct, discount_type, discount_val, shipping) -> dict:
    subtotal = q2(sum((i["amount"] for i in items), Decimal("0")))
    dval = to_dec(discount_val)
    if discount_type == "Percentage (%)":
        discount = q2(subtotal * min(dval, Decimal("100")) / Decimal("100"))
    elif discount_type == "Fixed amount":
        discount = q2(min(dval, subtotal))
    else:
        discount = Decimal("0.00")
    taxable = subtotal - discount
    tax_rate = to_dec(tax_pct)
    tax = q2(taxable * tax_rate / Decimal("100"))
    ship = q2(to_dec(shipping))
    total = q2(taxable + tax + ship)
    return {
        "subtotal": subtotal, "discount": discount, "taxable": taxable,
        "tax_rate": tax_rate, "tax": tax, "shipping": ship, "total": total,
    }


def discount_label(inv) -> str:
    if inv["discount_type"] == "Percentage (%)":
        return f"Discount ({fmt_qty(to_dec(inv['discount_val']))}%)"
    return "Discount"


# --------------------------------------------------------------------------
# Logo helper
# --------------------------------------------------------------------------
def normalize_logo(raw: bytes):
    """Convert any uploaded image to a reasonably small PNG (or None if invalid)."""
    if not raw:
        return None
    try:
        img = PILImage.open(io.BytesIO(raw))
        img = img.convert("RGBA")
        img.thumbnail((500, 500))
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()
    except Exception:
        return None


# --------------------------------------------------------------------------
# HTML preview (also printable)
# --------------------------------------------------------------------------
INVOICE_CSS = """
*{box-sizing:border-box}
body{margin:0;background:#f1f5f9;font-family:'Segoe UI',Roboto,Helvetica,Arial,sans-serif;color:#1e293b}
.wrap{padding:12px}
.toolbar{display:flex;justify-content:flex-end;margin-bottom:10px}
.btn{background:__ACCENT__;color:#fff;border:0;padding:8px 16px;border-radius:8px;font-size:13px;cursor:pointer;font-weight:600}
.btn:hover{opacity:.9}
.paper{background:#fff;border-radius:12px;box-shadow:0 4px 24px rgba(15,23,42,.10);padding:36px 38px;max-width:820px;margin:0 auto;border-top:6px solid __ACCENT__}
.head{display:flex;justify-content:space-between;gap:20px;flex-wrap:wrap}
.logo{max-height:70px;max-width:200px;margin-bottom:8px;display:block}
.biz{font-size:20px;font-weight:700;color:#0f172a}
.muted{color:#64748b;font-size:12.5px;line-height:1.55}
.title{font-size:34px;font-weight:800;letter-spacing:2px;color:__ACCENT__;text-align:right}
.meta{text-align:right;font-size:13px;line-height:1.7}
.meta b{color:#0f172a}
.billto{margin:26px 0 18px}
.label{font-size:11px;font-weight:700;letter-spacing:1.5px;color:__ACCENT__;margin-bottom:4px}
.cust{font-weight:600;font-size:15px}
.tbl-wrap{overflow-x:auto}
table.items{width:100%;border-collapse:collapse;font-size:13.5px;min-width:480px}
table.items th{background:__ACCENT__;color:#fff;text-align:left;padding:10px 12px;font-weight:600}
table.items td{padding:10px 12px;border-bottom:1px solid #e2e8f0;vertical-align:top}
table.items tr:nth-child(even) td{background:#f8fafc}
.r{text-align:right !important}
.totals{margin-left:auto;margin-top:16px;width:300px;max-width:100%;font-size:14px}
.totals div{display:flex;justify-content:space-between;padding:6px 12px}
.totals .grand{background:__ACCENT__;color:#fff;font-weight:700;font-size:16px;border-radius:8px;margin-top:6px;padding:10px 12px}
.notes{margin-top:28px;font-size:12.5px;color:#475569;line-height:1.6}
.foot{margin-top:28px;text-align:center;color:#94a3b8;font-size:12px;border-top:1px dashed #cbd5e1;padding-top:12px}
.empty{text-align:center;color:#94a3b8;padding:24px}
@media print{
 body{background:#fff}.toolbar{display:none}.wrap{padding:0}
 .paper{box-shadow:none;border-radius:0;max-width:100%;padding:10px 0}
 *{-webkit-print-color-adjust:exact;print-color-adjust:exact}
}
"""


def br(text: str) -> str:
    return html_escape(text or "").replace("\n", "<br>")


def build_invoice_html(inv, logo_png, items, t) -> str:
    sym = CURRENCIES.get(inv["currency"], "")
    css = INVOICE_CSS.replace("__ACCENT__", inv["accent"])

    logo_html = ""
    if logo_png:
        b64 = base64.b64encode(logo_png).decode()
        logo_html = f'<img class="logo" src="data:image/png;base64,{b64}" alt="logo">'

    biz_lines = []
    if inv["biz_address"]:
        biz_lines.append(br(inv["biz_address"]))
    contact = " &nbsp;•&nbsp; ".join(html_escape(x) for x in (inv["biz_email"], inv["biz_phone"]) if x)
    if contact:
        biz_lines.append(contact)
    if inv["biz_tax_id"]:
        biz_lines.append("Tax ID: " + html_escape(inv["biz_tax_id"]))

    cust_lines = []
    if inv["cust_address"]:
        cust_lines.append(br(inv["cust_address"]))
    contact_c = " &nbsp;•&nbsp; ".join(html_escape(x) for x in (inv["cust_email"], inv["cust_phone"]) if x)
    if contact_c:
        cust_lines.append(contact_c)

    if items:
        rows = ""
        for n, it in enumerate(items, 1):
            rows += (
                f"<tr><td>{n}</td><td>{html_escape(it['desc'])}</td>"
                f"<td class='r'>{fmt_qty(it['qty'])}</td>"
                f"<td class='r'>{money(sym, it['price'])}</td>"
                f"<td class='r'>{money(sym, it['amount'])}</td></tr>"
            )
        table = (
            "<div class='tbl-wrap'><table class='items'><thead><tr><th>#</th><th>Description</th>"
            "<th class='r'>Qty</th><th class='r'>Unit Price</th><th class='r'>Amount</th></tr></thead>"
            f"<tbody>{rows}</tbody></table></div>"
        )
    else:
        table = "<div class='empty'>Add at least one item to see it here.</div>"

    totals = f"<div><span>Subtotal</span><span>{money(sym, t['subtotal'])}</span></div>"
    if t["discount"] > 0:
        totals += f"<div><span>{html_escape(discount_label(inv))}</span><span>-{money(sym, t['discount'])}</span></div>"
    totals += f"<div><span>Tax ({fmt_qty(t['tax_rate'])}%)</span><span>{money(sym, t['tax'])}</span></div>"
    if t["shipping"] > 0:
        totals += f"<div><span>Shipping</span><span>{money(sym, t['shipping'])}</span></div>"
    totals += f"<div class='grand'><span>Total Due</span><span>{money(sym, t['total'])}</span></div>"

    notes = ""
    if inv["notes"]:
        notes += f"<div class='label'>NOTES</div><div>{br(inv['notes'])}</div><br>"
    if inv["terms"]:
        notes += f"<div class='label'>TERMS &amp; CONDITIONS</div><div>{br(inv['terms'])}</div>"

    biz_name = html_escape(inv["biz_name"]) or "Your Business Name"
    cust_name = html_escape(inv["cust_name"]) or "Customer Name"
    due = f"<br>Due Date: <b>{inv['due_date']:%d %b %Y}</b>" if inv["due_date"] else ""

    return f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><style>{css}</style></head>
<body><div class="wrap">
<div class="toolbar"><button class="btn" onclick="window.print()">🖨️ Print</button></div>
<div class="paper">
 <div class="head">
  <div>{logo_html}<div class="biz">{biz_name}</div><div class="muted">{'<br>'.join(biz_lines)}</div></div>
  <div><div class="title">INVOICE</div>
   <div class="meta">Invoice No: <b>{html_escape(inv['inv_no'])}</b><br>
   Date: <b>{inv['inv_date']:%d %b %Y}</b>{due}</div></div>
 </div>
 <div class="billto"><div class="label">BILL TO</div><div class="cust">{cust_name}</div>
  <div class="muted">{'<br>'.join(cust_lines)}</div></div>
 {table}
 <div class="totals">{totals}</div>
 <div class="notes">{notes}</div>
 <div class="foot">Thank you for your business!</div>
</div></div></body></html>"""


# --------------------------------------------------------------------------
# PDF export (ReportLab)
# --------------------------------------------------------------------------
def make_pdf(inv, logo_png, items, t) -> bytes:
    sym = CURRENCIES.get(inv["currency"], "")
    accent = colors.HexColor(inv["accent"])
    ink = colors.HexColor("#1e293b")
    muted = colors.HexColor("#64748b")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=18 * mm,
        title=f"Invoice {inv['inv_no']}", author=inv["biz_name"] or "Invoice Generator",
    )
    usable = A4[0] - 36 * mm

    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9.5, leading=13.5, textColor=ink)
    small = ParagraphStyle("small", parent=base, fontSize=9, leading=12.5, textColor=muted)
    biz_style = ParagraphStyle("biz", parent=base, fontName="Helvetica-Bold", fontSize=16, leading=20)
    title_style = ParagraphStyle("title", parent=base, fontName="Helvetica-Bold", fontSize=26,
                                 leading=30, textColor=accent, alignment=TA_RIGHT)
    meta_style = ParagraphStyle("meta", parent=base, alignment=TA_RIGHT, leading=16)
    label_style = ParagraphStyle("label", parent=base, fontName="Helvetica-Bold", fontSize=8.5,
                                 textColor=accent, spaceAfter=2)
    cust_style = ParagraphStyle("cust", parent=base, fontName="Helvetica-Bold", fontSize=11.5, leading=15)

    def para(text, style):
        return Paragraph(xml_escape(text or "").replace("\n", "<br/>"), style)

    # ---- header
    left = []
    if logo_png:
        w, h = PILImage.open(io.BytesIO(logo_png)).size
        scale = min(48 * mm / w, 20 * mm / h)
        img = RLImage(io.BytesIO(logo_png), width=w * scale, height=h * scale)
        img.hAlign = "LEFT"
        left += [img, Spacer(1, 3 * mm)]
    left.append(para(inv["biz_name"] or "Your Business Name", biz_style))
    details = []
    if inv["biz_address"]:
        details.append(inv["biz_address"])
    contact = "  |  ".join(x for x in (inv["biz_email"], inv["biz_phone"]) if x)
    if contact:
        details.append(contact)
    if inv["biz_tax_id"]:
        details.append("Tax ID: " + inv["biz_tax_id"])
    if details:
        left.append(para("\n".join(details), small))

    meta = f"Invoice No: <b>{xml_escape(inv['inv_no'])}</b><br/>Date: <b>{inv['inv_date']:%d %b %Y}</b>"
    if inv["due_date"]:
        meta += f"<br/>Due Date: <b>{inv['due_date']:%d %b %Y}</b>"
    right = [Paragraph("INVOICE", title_style), Spacer(1, 2 * mm), Paragraph(meta, meta_style)]

    header = Table([[left, right]], colWidths=[usable * 0.58, usable * 0.42])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))

    story = [header, Spacer(1, 4 * mm), HRFlowable(width="100%", thickness=2, color=accent), Spacer(1, 5 * mm)]

    # ---- bill to
    story.append(Paragraph("BILL TO", label_style))
    story.append(para(inv["cust_name"] or "Customer Name", cust_style))
    cdetails = []
    if inv["cust_address"]:
        cdetails.append(inv["cust_address"])
    contact_c = "  |  ".join(x for x in (inv["cust_email"], inv["cust_phone"]) if x)
    if contact_c:
        cdetails.append(contact_c)
    if cdetails:
        story.append(para("\n".join(cdetails), small))
    story.append(Spacer(1, 7 * mm))

    # ---- items
    desc_style = ParagraphStyle("desc", parent=base, fontSize=9.5, leading=12.5)
    data = [["#", "Description", "Qty", "Unit Price", "Amount"]]
    for n, it in enumerate(items, 1):
        data.append([str(n), para(it["desc"], desc_style), fmt_qty(it["qty"]),
                     money(sym, it["price"]), money(sym, it["amount"])])
    if not items:
        data.append(["", "No items", "", "", ""])
    tbl = Table(data, colWidths=[10 * mm, usable - 10 * mm - 20 * mm - 32 * mm - 32 * mm,
                                 20 * mm, 32 * mm, 32 * mm], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), accent),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("TEXTCOLOR", (0, 1), (-1, -1), ink),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEBELOW", (0, 1), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
    ]
    for r in range(2, len(data), 2):
        style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#f8fafc")))
    tbl.setStyle(TableStyle(style))
    story += [tbl, Spacer(1, 5 * mm)]

    # ---- totals
    rows = [["Subtotal", money(sym, t["subtotal"])]]
    if t["discount"] > 0:
        rows.append([discount_label(inv), "-" + money(sym, t["discount"])])
    rows.append([f"Tax ({fmt_qty(t['tax_rate'])}%)", money(sym, t["tax"])])
    if t["shipping"] > 0:
        rows.append(["Shipping", money(sym, t["shipping"])])
    rows.append(["Total Due", money(sym, t["total"])])
    tot = Table(rows, colWidths=[42 * mm, 38 * mm], hAlign="RIGHT")
    tot.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TEXTCOLOR", (0, 0), (-1, -1), ink),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("BACKGROUND", (0, -1), (-1, -1), accent),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, -1), (-1, -1), 11.5),
        ("TOPPADDING", (0, -1), (-1, -1), 8),
        ("BOTTOMPADDING", (0, -1), (-1, -1), 8),
    ]))
    story.append(tot)

    # ---- notes
    if inv["notes"]:
        story += [Spacer(1, 9 * mm), Paragraph("NOTES", label_style), para(inv["notes"], small)]
    if inv["terms"]:
        story += [Spacer(1, 5 * mm), Paragraph("TERMS &amp; CONDITIONS", label_style), para(inv["terms"], small)]

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8.5)
        canvas.setFillColor(muted)
        canvas.drawCentredString(A4[0] / 2, 10 * mm, f"Thank you for your business!   -   Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Excel export (openpyxl) - uses real formulas so the sheet stays editable
# --------------------------------------------------------------------------
def make_excel(inv, logo_png, items, t) -> bytes:
    sym = CURRENCIES.get(inv["currency"], "")
    accent = inv["accent"].lstrip("#").upper()
    money_fmt = f'"{sym}"#,##0.00'

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (6, 46, 12, 18, 20)):
        ws.column_dimensions[col].width = w

    white_bold = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color=accent, end_color=accent)
    thin = Side(style="thin", color="E2E8F0")

    # logo occupies rows 1-4
    if logo_png:
        try:
            img = XLImage(io.BytesIO(logo_png))
            scale = min(190 / img.width, 70 / img.height, 1)
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            ws.add_image(img, "A1")
        except Exception:
            pass
    for r in range(1, 5):
        ws.row_dimensions[r].height = 18

    ws["A5"] = inv["biz_name"] or "Your Business Name"
    ws["A5"].font = Font(bold=True, size=16)
    ws.merge_cells("A5:C5")
    ws["D5"] = "INVOICE"
    ws["D5"].font = Font(bold=True, size=20, color=accent)
    ws["D5"].alignment = Alignment(horizontal="right")
    ws.merge_cells("D5:E5")

    ws["A6"] = inv["biz_address"]
    ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A6:C6")
    ws.row_dimensions[6].height = 42
    ws["A7"] = "  |  ".join(x for x in (inv["biz_email"], inv["biz_phone"]) if x)
    ws.merge_cells("A7:C7")
    ws["A8"] = ("Tax ID: " + inv["biz_tax_id"]) if inv["biz_tax_id"] else ""
    ws.merge_cells("A8:C8")

    meta = [("Invoice No:", inv["inv_no"]), ("Date:", inv["inv_date"]), ("Due Date:", inv["due_date"])]
    for i, (lab, val) in enumerate(meta):
        r = 6 + i
        ws.cell(r, 4, lab).font = Font(bold=True)
        ws.cell(r, 4).alignment = Alignment(horizontal="right")
        c = ws.cell(r, 5, val)
        c.alignment = Alignment(horizontal="right")
        if isinstance(val, date):
            c.number_format = "DD MMM YYYY"

    ws["A10"] = "BILL TO"
    ws["A10"].font = Font(bold=True, color=accent, size=9)
    ws["A11"] = inv["cust_name"] or "Customer Name"
    ws["A11"].font = Font(bold=True, size=12)
    ws["A12"] = inv["cust_address"]
    ws["A12"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A12:C12")
    ws.row_dimensions[12].height = 42
    ws["A13"] = "  |  ".join(x for x in (inv["cust_email"], inv["cust_phone"]) if x)
    ws.merge_cells("A13:C13")

    # ---- items table
    hr = 15
    for c, h in enumerate(["#", "Description", "Qty", "Unit Price", "Amount"], 1):
        cell = ws.cell(hr, c, h)
        cell.font = white_bold
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="left" if c <= 2 else "right", vertical="center")
    ws.row_dimensions[hr].height = 22

    first = hr + 1
    r = first
    for n, it in enumerate(items, 1):
        ws.cell(r, 1, n).alignment = Alignment(horizontal="left", vertical="top")
        d = ws.cell(r, 2, it["desc"])
        d.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 3, float(it["qty"])).number_format = "General"
        ws.cell(r, 4, float(it["price"])).number_format = money_fmt
        ws.cell(r, 5, f"=ROUND(C{r}*D{r},2)").number_format = money_fmt
        for c in range(1, 6):
            ws.cell(r, c).border = Border(bottom=thin)
            if c >= 3:
                ws.cell(r, c).alignment = Alignment(horizontal="right", vertical="top")
        r += 1
    last = max(r - 1, first)

    # ---- totals (formulas)
    r += 1
    sub_r = r
    ws.cell(sub_r, 4, "Subtotal").font = Font(bold=True)
    ws.cell(sub_r, 5, f"=SUM(E{first}:E{last})" if items else 0)

    r += 1
    disc_r = r
    ws.cell(disc_r, 4, "Discount")
    if inv["discount_type"] == "Percentage (%)":
        ws.cell(disc_r, 3, float(min(to_dec(inv["discount_val"]), Decimal("100"))))
        ws.cell(disc_r, 3).number_format = '0.##"%"'
        ws.cell(disc_r, 5, f"=-ROUND(E{sub_r}*C{disc_r}/100,2)")
    elif inv["discount_type"] == "Fixed amount":
        ws.cell(disc_r, 3, float(to_dec(inv["discount_val"])))
        ws.cell(disc_r, 3).number_format = money_fmt
        ws.cell(disc_r, 5, f"=-MIN(C{disc_r},E{sub_r})")
    else:
        ws.cell(disc_r, 5, 0)

    r += 1
    tax_r = r
    ws.cell(tax_r, 4, "Tax")
    ws.cell(tax_r, 3, float(t["tax_rate"])).number_format = '0.##"%"'
    ws.cell(tax_r, 5, f"=ROUND((E{sub_r}+E{disc_r})*C{tax_r}/100,2)")

    r += 1
    ship_r = r
    ws.cell(ship_r, 4, "Shipping")
    ws.cell(ship_r, 5, float(t["shipping"]))

    r += 1
    tot_r = r
    ws.cell(tot_r, 4, "TOTAL DUE")
    ws.cell(tot_r, 5, f"=E{sub_r}+E{disc_r}+E{tax_r}+E{ship_r}")

    for rr in (sub_r, disc_r, tax_r, ship_r, tot_r):
        ws.cell(rr, 4).alignment = Alignment(horizontal="right")
        ws.cell(rr, 5).number_format = money_fmt
        ws.cell(rr, 5).alignment = Alignment(horizontal="right")
    for rr in (disc_r, tax_r):
        ws.cell(rr, 3).font = Font(color="2563EB")  # blue = editable input
        ws.cell(rr, 3).alignment = Alignment(horizontal="right")
    for c in (4, 5):
        cell = ws.cell(tot_r, c)
        cell.font = white_bold
        cell.fill = head_fill

    r = tot_r + 2
    if inv["notes"]:
        ws.cell(r, 1, "NOTES").font = Font(bold=True, color=accent, size=9)
        r += 1
        ws.cell(r, 1, inv["notes"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 15 * max(2, inv["notes"].count("\n") + 2)
        r += 2
    if inv["terms"]:
        ws.cell(r, 1, "TERMS & CONDITIONS").font = Font(bold=True, color=accent, size=9)
        r += 1
        ws.cell(r, 1, inv["terms"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 15 * max(2, inv["terms"].count("\n") + 2)

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------
# Gemini helpers (plain REST call - no SDK needed, free tier key is enough)
# --------------------------------------------------------------------------
def call_gemini(prompt, api_key, model=DEFAULT_MODEL, json_mode=False, system=None) -> str:
    """Call Gemini via REST. Falls back to other Flash models if one is missing (404) or out of quota (429)."""
    body = {"contents": [{"parts": [{"text": prompt}]}], "generationConfig": {"temperature": 0.3}}
    if json_mode:
        body["generationConfig"]["responseMimeType"] = "application/json"
    if system:
        body["systemInstruction"] = {"parts": [{"text": system}]}

    candidates = [model] + [m for m in FALLBACK_MODELS if m != model]
    last_error = "Unknown error"
    for name in candidates:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{name}:generateContent"
        try:
            resp = requests.post(
                url, json=body, timeout=60,
                headers={"Content-Type": "application/json", "x-goog-api-key": api_key},
            )
        except requests.RequestException as exc:
            raise RuntimeError(f"Network error: {exc}") from exc
        if resp.status_code == 200:
            try:
                parts = resp.json()["candidates"][0]["content"]["parts"]
                return "".join(p.get("text", "") for p in parts).strip()
            except (KeyError, IndexError, ValueError) as exc:
                raise RuntimeError("Gemini returned an empty or blocked response. Try rephrasing.") from exc
        try:
            msg = resp.json().get("error", {}).get("message", resp.text)
        except Exception:
            msg = resp.text
        last_error = f"Gemini API error ({resp.status_code}) on '{name}': {msg[:300]}"
        if resp.status_code not in (404, 429):
            break  # bad key, bad request, etc. - another model will not help
    raise RuntimeError(last_error)


def parse_json_loose(text: str):
    """Parse JSON even if the model wrapped it in ```json fences or added chatter."""
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip(), flags=re.I)
    try:
        return json.loads(text)
    except ValueError:
        pass
    m = re.search(r"(\{.*\}|\[.*\])", text, flags=re.S)
    if m:
        return json.loads(m.group(1))
    raise ValueError("Could not read the AI response as JSON.")


def get_api_key() -> str:
    key = (st.session_state.get("api_key_input") or "").strip()
    if key:
        return key
    try:
        key = str(st.secrets.get("GEMINI_API_KEY", "")).strip()
    except Exception:
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "").strip()


# --------------------------------------------------------------------------
# State management
# --------------------------------------------------------------------------
def blank_form() -> dict:
    today = date.today()
    return {
        "biz_name": "", "biz_email": "", "biz_phone": "", "biz_address": "", "biz_tax_id": "",
        "cust_name": "", "cust_email": "", "cust_phone": "", "cust_address": "",
        "inv_no": f"INV-{today:%Y%m%d}-001", "inv_date": today, "due_date": today + timedelta(days=14),
        "currency": "USD", "accent": "#4F46E5", "notes": "", "terms": "",
        "tax_pct": 0.0, "discount_type": "None", "discount_val": 0.0, "shipping": 0.0,
    }


def blank_items() -> pd.DataFrame:
    return pd.DataFrame({"Description": [""], "Quantity": [1.0], "Unit Price": [0.0]})


def sample_form() -> dict:
    f = blank_form()
    f.update({
        "biz_name": "Bright Studio Co.", "biz_email": "hello@brightstudio.com", "biz_phone": "+1 555 010 2020",
        "biz_address": "21 Market Street\nSpringfield, 12345", "biz_tax_id": "TAX-123456",
        "cust_name": "Acme Corporation", "cust_email": "accounts@acme.com", "cust_phone": "+1 555 777 8888",
        "cust_address": "500 Industrial Ave\nMetropolis, 54321",
        "notes": "Thank you for choosing Bright Studio. Payment is due within 14 days.",
        "terms": "Bank transfer preferred. Late payments may incur a 2% monthly fee.",
        "tax_pct": 10.0, "discount_type": "Percentage (%)", "discount_val": 5.0, "shipping": 0.0,
    })
    return f


def sample_items() -> pd.DataFrame:
    return pd.DataFrame({
        "Description": ["Brand identity design", "Website development (hours)", "Hosting - 12 months"],
        "Quantity": [1.0, 20.0, 12.0],
        "Unit Price": [450.00, 55.00, 9.99],
    })


def _apply_state(form: dict, items: pd.DataFrame):
    for k, v in form.items():
        st.session_state[k] = v
    st.session_state["items_base"] = items
    st.session_state["items_ver"] = st.session_state.get("items_ver", 0) + 1
    st.session_state["logo_ver"] = st.session_state.get("logo_ver", 0) + 1


def init_state():
    if "initialized" not in st.session_state:
        _apply_state(sample_form(), sample_items())
        st.session_state["initialized"] = True
        st.session_state.setdefault("model_name", DEFAULT_MODEL)


def reset_form():  # callback
    _apply_state(blank_form(), blank_items())
    st.session_state["ai_msg"] = None


def load_sample():  # callback
    _apply_state(sample_form(), sample_items())


# --------------------------------------------------------------------------
# AI jobs (run at the top of the script, before widgets are created)
# --------------------------------------------------------------------------
def run_pending_ai():
    job = st.session_state.pop("ai_job", None)
    if not job:
        return
    key = get_api_key()
    if not key:
        st.session_state["ai_msg"] = ("error", "Please add your free Gemini API key in the sidebar first.")
        return
    model = (st.session_state.get("model_name") or DEFAULT_MODEL).strip()
    kind = job["kind"]
    try:
        with st.spinner("Gemini is working..."):
            if kind == "smart_fill":
                _job_smart_fill(job["text"], key, model)
            elif kind == "notes":
                _job_notes(job, key, model)
            elif kind == "polish":
                _job_polish(key, model)
    except Exception as exc:  # show any problem nicely instead of crashing
        st.session_state["ai_msg"] = ("error", str(exc))


def _job_smart_fill(text, key, model):
    prompt = (
        "Extract invoice data from the text below. Reply with ONLY JSON in this exact shape:\n"
        '{"customer_name": str|null, "customer_email": str|null, "customer_address": str|null, '
        '"tax_percent": number|null, "items": [{"description": str, "quantity": number, "unit_price": number}]}\n'
        "Use quantity 1 when not stated. Prices are per unit, as plain numbers without currency symbols.\n\n"
        f"TEXT:\n{text}"
    )
    data = parse_json_loose(call_gemini(prompt, key, model, json_mode=True))
    if isinstance(data, list):
        data = {"items": data}
    new_rows = []
    for it in data.get("items") or []:
        new_rows.append({
            "Description": str(it.get("description", "")).strip(),
            "Quantity": float(to_dec(it.get("quantity", 1)) or 1),
            "Unit Price": float(to_dec(it.get("unit_price", 0))),
        })
    if not new_rows:
        raise RuntimeError("No items were found in that text. Try adding quantities and prices.")
    current = st.session_state.get("items_current", blank_items())
    keep = clean_items(current)
    base = pd.DataFrame(
        [{"Description": i["desc"], "Quantity": float(i["qty"]), "Unit Price": float(i["price"])} for i in keep],
        columns=ITEM_COLS,
    )
    st.session_state["items_base"] = pd.concat([base, pd.DataFrame(new_rows)], ignore_index=True)
    st.session_state["items_ver"] += 1
    for src, dst in (("customer_name", "cust_name"), ("customer_email", "cust_email"),
                     ("customer_address", "cust_address")):
        if data.get(src):
            st.session_state[dst] = str(data[src])
    if data.get("tax_percent") not in (None, ""):
        st.session_state["tax_pct"] = float(min(to_dec(data["tax_percent"]), Decimal("100")))
    st.session_state["ai_msg"] = ("success", f"Added {len(new_rows)} item(s) from your text. Check them in the Items tab.")


def _job_notes(job, key, model):
    inv = job["context"]
    ss = st.session_state
    t = compute_totals(clean_items(ss.get("items_current", blank_items())), ss.get("tax_pct"),
                       ss.get("discount_type"), ss.get("discount_val"), ss.get("shipping"))
    total = money(CURRENCIES.get(ss.get("currency"), ""), t["total"])
    prompt = (
        f"Write a short, {job['tone'].lower()} invoice note (max 60 words) from '{inv['biz']}' to '{inv['cust']}'. "
        f"Invoice total is {total}, due {inv['due']}. Mention the payment due date. "
        "Plain text only: no markdown, no subject line, no placeholders in brackets."
    )
    st.session_state["notes"] = call_gemini(prompt, key, model)
    if job.get("with_terms"):
        terms_prompt = (
            "Write 3 short, standard payment terms for a small business invoice (late fee, accepted payment "
            "methods, dispute window). Plain text, one per line, no numbering or markdown."
        )
        st.session_state["terms"] = call_gemini(terms_prompt, key, model)
    st.session_state["ai_msg"] = ("success", "Notes written. Review them in the Invoice tab.")


def _job_polish(key, model):
    items = clean_items(st.session_state.get("items_current", blank_items()))
    if not items:
        raise RuntimeError("Add some items first.")
    prompt = (
        "Rewrite each invoice line-item description to be clear, professional and concise (max 10 words each). "
        "Keep the meaning. Reply with ONLY a JSON array of strings, same order and same length as the input.\n"
        + json.dumps([i["desc"] for i in items])
    )
    out = parse_json_loose(call_gemini(prompt, key, model, json_mode=True))
    if not isinstance(out, list) or len(out) != len(items):
        raise RuntimeError("The AI response didn't match your items. Please try again.")
    st.session_state["items_base"] = pd.DataFrame(
        [{"Description": str(new), "Quantity": float(i["qty"]), "Unit Price": float(i["price"])}
         for new, i in zip(out, items)],
        columns=ITEM_COLS,
    )
    st.session_state["items_ver"] += 1
    st.session_state["ai_msg"] = ("success", "Descriptions polished.")


# --------------------------------------------------------------------------
# Styling
# --------------------------------------------------------------------------
APP_CSS = """
<style>
.block-container{padding-top:1.4rem;max-width:1400px}
.hero{background:linear-gradient(120deg,#4f46e5 0%,#7c3aed 55%,#06b6d4 100%);color:#fff;
      border-radius:16px;padding:22px 28px;margin-bottom:18px}
.hero h1{margin:0;font-size:1.9rem;color:#fff;padding:0}
.hero p{margin:4px 0 0;opacity:.92}
.kpi{background:var(--secondary-background-color);border-radius:14px;padding:14px 18px;
     border-left:5px solid #4f46e5}
.kpi .lab{font-size:.78rem;letter-spacing:.06em;text-transform:uppercase;opacity:.7}
.kpi .val{font-size:1.45rem;font-weight:700;margin-top:2px}
.kpi.green{border-left-color:#10b981}.kpi.amber{border-left-color:#f59e0b}.kpi.cyan{border-left-color:#06b6d4}
div[data-testid="stDownloadButton"] button{width:100%;font-weight:600}
</style>
"""


def show_html(html: str, height: int):
    """Render an HTML string. Uses st.iframe on new Streamlit, falls back for older versions."""
    if hasattr(st, "iframe"):
        st.iframe(html, height=height)
    else:  # Streamlit < 1.56
        import streamlit.components.v1 as components
        components.html(html, height=height, scrolling=True)


def kpi(label, value, cls=""):
    return f'<div class="kpi {cls}"><div class="lab">{label}</div><div class="val">{value}</div></div>'


# --------------------------------------------------------------------------
# Main UI
# --------------------------------------------------------------------------
def main():
    st.set_page_config(page_title="AI Invoice Generator", page_icon="🧾", layout="wide")
    st.markdown(APP_CSS, unsafe_allow_html=True)
    init_state()
    run_pending_ai()

    # ---------------- sidebar
    with st.sidebar:
        st.header("⚙️ Control Panel")
        st.button("🔄 Reset form", on_click=reset_form, width="stretch")
        st.button("🧪 Load sample data", on_click=load_sample, width="stretch")
        st.divider()
        st.subheader("✨ Gemini AI (optional)")
        st.text_input("Gemini API key", type="password", key="api_key_input",
                      help="Free key from https://aistudio.google.com/apikey. Never stored or shared.")
        st.text_input("Model", key="model_name", help="Default \"gemini-flash-latest\" always points to Google's newest Flash model.")
        if get_api_key():
            st.success("AI features enabled")
        else:
            st.info("Everything else works without a key.")
        st.caption("Tip: on Streamlit Cloud, put GEMINI_API_KEY in app Secrets.")

    st.markdown(
        '<div class="hero"><h1>🧾 AI Invoice Generator</h1>'
        "<p>Create professional invoices with a live preview, then download as PDF or Excel.</p></div>",
        unsafe_allow_html=True,
    )
    kpi_box = st.container()

    left, right = st.columns([1.05, 1], gap="large")

    # ---------------- left: form
    with left:
        t_biz, t_cust, t_inv, t_items, t_ai = st.tabs(
            ["🏢 Business", "👤 Customer", "🧾 Invoice", "📦 Items & Tax", "✨ AI Assistant"]
        )

        with t_biz:
            st.text_input("Business name", key="biz_name", placeholder="e.g. Bright Studio Co.")
            up = st.file_uploader("Logo (PNG / JPG / WEBP)", type=["png", "jpg", "jpeg", "webp"],
                                  key=f"logo_{st.session_state.get('logo_ver', 0)}")
            c1, c2 = st.columns(2)
            c1.text_input("Email", key="biz_email")
            c2.text_input("Phone", key="biz_phone")
            st.text_area("Address", key="biz_address", height=80)
            st.text_input("Tax / VAT / NTN ID (optional)", key="biz_tax_id")

        with t_cust:
            st.text_input("Customer name", key="cust_name", placeholder="e.g. Acme Corporation")
            c1, c2 = st.columns(2)
            c1.text_input("Customer email", key="cust_email")
            c2.text_input("Customer phone", key="cust_phone")
            st.text_area("Billing address", key="cust_address", height=80)

        with t_inv:
            c1, c2 = st.columns(2)
            c1.text_input("Invoice number", key="inv_no")
            c2.selectbox("Currency", list(CURRENCIES), key="currency")
            c3, c4 = st.columns(2)
            c3.date_input("Invoice date", key="inv_date")
            c4.date_input("Due date", key="due_date")
            if st.session_state["due_date"] < st.session_state["inv_date"]:
                st.warning("Due date is before the invoice date.")
            st.color_picker("Brand / accent colour", key="accent")
            st.text_area("Notes", key="notes", height=90)
            st.text_area("Terms & conditions", key="terms", height=90)

        with t_items:
            st.caption("Click the empty bottom row to add a product. Select a row and press Delete to remove it.")
            edited = st.data_editor(
                st.session_state["items_base"],
                key=f"editor_{st.session_state['items_ver']}",
                num_rows="dynamic",
                width="stretch",
                hide_index=True,
                column_config={
                    "Description": st.column_config.TextColumn("Product / Service", width="large"),
                    "Quantity": st.column_config.NumberColumn("Qty", min_value=0.0, step=1.0, default=1.0, format="%g"),
                    "Unit Price": st.column_config.NumberColumn("Unit Price", min_value=0.0, step=0.5, default=0.0, format="%.2f"),
                },
            )
            st.session_state["items_current"] = edited
            c1, c2 = st.columns(2)
            c1.number_input("Tax (%)", min_value=0.0, max_value=100.0, step=0.5, key="tax_pct")
            c2.number_input("Shipping / extra fees", min_value=0.0, step=1.0, key="shipping")
            c3, c4 = st.columns(2)
            c3.selectbox("Discount", DISCOUNT_TYPES, key="discount_type")
            c4.number_input("Discount value", min_value=0.0, step=1.0, key="discount_val",
                            disabled=st.session_state["discount_type"] == "None")

        with t_ai:
            st.markdown("**1. Smart fill** - paste an email, chat or rough list and Gemini builds the line items.")
            st.text_area("Describe the work", key="ai_text", height=110,
                         placeholder="e.g. Acme Corp ordered 3 logo designs at 150 each and 2 hours of consulting at 80. Tax 8%.")
            if st.button("✨ Extract items", key="btn_fill"):
                if (st.session_state.get("ai_text") or "").strip():
                    st.session_state["ai_job"] = {"kind": "smart_fill", "text": st.session_state["ai_text"]}
                    st.rerun()
                else:
                    st.session_state["ai_msg"] = ("error", "Please type or paste some text first.")
            st.divider()
            st.markdown("**2. Write notes & terms**")
            tone = st.selectbox("Tone", ["Friendly", "Professional", "Formal", "Firm (overdue)"], key="ai_tone")
            with_terms = st.checkbox("Also write payment terms", value=True, key="ai_terms")
            if st.button("📝 Write notes", key="btn_notes"):
                st.session_state["ai_job"] = {
                    "kind": "notes", "tone": tone, "with_terms": with_terms,
                    "context": {
                        "biz": st.session_state["biz_name"] or "our company",
                        "cust": st.session_state["cust_name"] or "the customer",
                        "due": f"{st.session_state['due_date']:%d %b %Y}",
                    },
                }
                st.rerun()
            st.divider()
            st.markdown("**3. Polish item descriptions**")
            if st.button("🪄 Polish descriptions", key="btn_polish"):
                st.session_state["ai_job"] = {"kind": "polish"}
                st.rerun()
            msg = st.session_state.get("ai_msg")
            if msg:
                (st.success if msg[0] == "success" else st.error)(msg[1])

    # ---------------- build invoice data
    ss = st.session_state
    inv = {k: ss.get(k) for k in FORM_KEYS}
    for k in ("biz_name", "biz_email", "biz_phone", "biz_address", "biz_tax_id",
              "cust_name", "cust_email", "cust_phone", "cust_address", "inv_no", "notes", "terms"):
        inv[k] = (inv[k] or "").strip()
    inv["inv_no"] = inv["inv_no"] or "INV-0001"
    logo_png = normalize_logo(up.getvalue()) if up is not None else None
    if up is not None and logo_png is None:
        st.sidebar.error("Could not read that logo image.")

    items = clean_items(edited)
    totals = compute_totals(items, inv["tax_pct"], inv["discount_type"], inv["discount_val"], inv["shipping"])
    sym = CURRENCIES.get(inv["currency"], "")

    # ---------------- KPI dashboard
    with kpi_box:
        k1, k2, k3, k4 = st.columns(4)
        k1.markdown(kpi("Subtotal", money(sym, totals["subtotal"])), unsafe_allow_html=True)
        k2.markdown(kpi(f"Tax ({fmt_qty(totals['tax_rate'])}%)", money(sym, totals["tax"]), "amber"), unsafe_allow_html=True)
        k3.markdown(kpi("Grand total", money(sym, totals["total"]), "green"), unsafe_allow_html=True)
        k4.markdown(kpi("Line items", str(len(items)), "cyan"), unsafe_allow_html=True)
        st.write("")

    # ---------------- right: preview + downloads
    with right:
        st.subheader("👁️ Live preview")
        html = build_invoice_html(inv, logo_png, items, totals)
        show_html(html, 780 + 42 * len(items))

        st.subheader("⬇️ Download")
        if not items:
            st.warning("Add at least one product to enable downloads.")
        else:
            if not inv["biz_name"] or not inv["cust_name"]:
                st.info("Tip: fill in the business and customer names for a complete invoice.")
            safe_no = re.sub(r"[^A-Za-z0-9_-]+", "_", inv["inv_no"])
            d1, d2 = st.columns(2)
            try:
                d1.download_button("📄 Download PDF", make_pdf(inv, logo_png, items, totals),
                                   file_name=f"invoice_{safe_no}.pdf", mime="application/pdf")
                d2.download_button("📊 Download Excel", make_excel(inv, logo_png, items, totals),
                                   file_name=f"invoice_{safe_no}.xlsx",
                                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            except Exception as exc:
                st.error(f"Could not generate the file: {exc}")


if __name__ == "__main__":
    main()
